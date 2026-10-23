#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""04_gerar_planilhas.py - Gera planilhas editaveis (com formulas) do dossie. Autor: Davi Fernandes Bonfim - Mat. 241003473."""
import pandas as pd, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
R = os.path.join(os.path.dirname(__file__), '..', 'resultados'); P = os.path.join(os.path.dirname(__file__), '..', 'planilhas')
os.makedirs(P, exist_ok=True)
AR = 'Arial'
f_norm = Font(name=AR, size=10); f_b = Font(name=AR, size=10, bold=True); f_in = Font(name=AR, size=10, color='0000FF')
f_h = Font(name=AR, size=10, bold=True, color='FFFFFF'); fill_h = PatternFill('solid', fgColor='1F3864'); fill_k = PatternFill('solid', fgColor='FFF2CC')
thin = Side(style='thin', color='BFBFBF'); bd = Border(top=thin, bottom=thin, left=thin, right=thin)
NUM = '#,##0;(#,##0);-'; NUM1 = '#,##0.0;(#,##0.0);-'; PCT = '0.00%'
def hdr(ws, r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = f_h; c.fill = fill_h; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = bd
def put(ws, r, c, v, fmt=None, inp=False, bold=False):
    x = ws.cell(r, c, v); x.font = f_in if inp else (f_b if bold else f_norm); x.border = bd
    if fmt: x.number_format = fmt
    return x
def title(ws, t, sub):
    ws['A1'] = t; ws['A1'].font = Font(name=AR, size=13, bold=True, color='1F3864'); ws['A2'] = sub; ws['A2'].font = Font(name=AR, size=9, italic=True)
def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1): ws.column_dimensions[L(i)].width = w
wb = Workbook()
# ---------------- Leia-me ----------------
ws = wb.active; ws.title = 'Leia-me'
title(ws, 'Planilhas de apoio ao dossie de auditoria - SABESP 2025', 'Davi Fernandes Bonfim - Matricula 241003473 - Auditoria 1 (UnB, CCA0024, Turma 02). Valores em R$ mil, salvo indicacao.')
txt = ['Fonte dos dados reais: DFP 2025 e 2024 da SABESP (CVM). Dados da base e da amostra sao SINTETICOS.',
       'Celulas em azul sao entradas (editaveis); as demais sao formulas. A coluna "Fonte" cita a pagina do PDF da DFP 2025 (F1).',
       'Abas: Materialidade (PT 320); Conciliacoes (PT 500); Analiticos (PT 520); Sensibilidade (PT 540); Amostragem (PT 530); Amostra (PT 530);',
       'Conc_Base (T-01); Erros (comparacao dos erros inseridos com os identificados); Distorcoes (PT 450).',
       'Premissas hardcoded estao identificadas ao lado de cada tabela. Nada aqui e opiniao de auditoria independente.']
for i, t in enumerate(txt, 4): ws.cell(i, 1, t).font = f_norm
widths(ws, [140])
# ---------------- Materialidade ----------------
ws = wb.create_sheet('Materialidade'); title(ws, 'PT 320 - Materialidade', 'R$ mil, DFP 2025 consolidada (F1, pp. 24-28)')
hdr(ws, 4, ['Base de referencia', 'Valor 2025', 'Percentual', 'Materialidade resultante', 'Papel', 'Fonte'])
rows = [('Lucro antes dos tributos (LAIR)', 11701851, 0.05, 'Adotado', 'F1, p. 28'), ('Receita liquida', 38092050, 0.01, 'Referencia', 'F1, p. 28'),
        ('Ativo total', 104202767, 0.005, 'Referencia', 'F1, p. 24'), ('Patrimonio liquido', 42401124, 0.01, 'Referencia', 'F1, p. 27')]
for i, (n, v, p, pa, fo) in enumerate(rows, 5):
    put(ws, i, 1, n); put(ws, i, 2, v, NUM, True); put(ws, i, 3, p, PCT, True); put(ws, i, 4, f'=B{i}*C{i}', NUM); put(ws, i, 5, pa); put(ws, i, 6, fo)
hdr(ws, 11, ['Parametro', 'Valor', 'Premissa', 'Calculo', '', ''])
put(ws, 12, 1, 'Materialidade global (MG)', bold=True); put(ws, 12, 2, '=ROUND(D5,0)', NUM, bold=True); put(ws, 12, 3, '5% do LAIR'); 
put(ws, 13, 1, 'Percentual da ME sobre a MG'); put(ws, 13, 2, 0.70, PCT, True); put(ws, 13, 3, 'Julgamento do aluno: auditoria inicial e riscos significativos')
put(ws, 14, 1, 'Materialidade para execucao (ME)', bold=True); put(ws, 14, 2, '=ROUND(B12*B13,0)', NUM, bold=True)
put(ws, 15, 1, 'Percentual do limite de distorcoes triviais'); put(ws, 15, 2, 0.05, PCT, True); put(ws, 15, 3, '5% da MG')
put(ws, 16, 1, 'Limite de distorcoes claramente triviais'); put(ws, 16, 2, '=ROUND(B12*B15,0)', NUM)
put(ws, 17, 1, 'Percentual da materialidade especifica (partes relacionadas)'); put(ws, 17, 2, 0.5, PCT, True)
put(ws, 18, 1, 'Materialidade especifica - partes relacionadas'); put(ws, 18, 2, '=ROUND(B12*B17,0)', NUM)
put(ws, 19, 1, 'Limite de investigacao analitica (25% da ME)'); put(ws, 19, 2, '=ROUND(B14*0.25,0)', NUM); put(ws, 19, 3, 'Premissa do aluno para riscos significativos')
put(ws, 20, 1, 'MG como % da receita'); put(ws, 20, 2, '=B12/B6', PCT); put(ws, 21, 1, 'MG como % do ativo'); put(ws, 21, 2, '=B12/B7', PCT)
widths(ws, [52, 16, 48, 24, 16, 14])
# ---------------- Conciliacoes ----------------
ws = wb.create_sheet('Conciliacoes'); title(ws, 'PT 500 - Conciliacoes (R$ mil, consolidado)', 'Entradas em azul vem das notas da DFP 2025 (F1); diferenca deve ser zero')
hdr(ws, 4, ['Conciliacao', 'Comp. 1', 'Comp. 2', 'Comp. 3', 'Comp. 4', 'Comp. 5', 'Comp. 6', 'Calculado', 'Saldo nas demonstracoes', 'Diferenca', 'Fonte'])
C = [('Intangivel: saldo inicial + movimentos', [44771124, 3207, 8117869, -139884, -98731, -2161112], 50492473, 'Nota 14, p. 73; balanco p. 25'),
     ('Intangivel: soma das classes', [137725, 75089, 3730121, 45946471, 551102, 51965], 50492473, 'Nota 14, p. 72'),
     ('Ativo de contrato: saldo inicial + movimentos', [4877667, 16684564, -62684, 24108, -8117869, -2387270], 11018516, 'Nota 13, p. 69; balanco p. 24'),
     ('Ativo financeiro: saldo inicial + movimentos', [17601626, 2387270, 1676434], 21665330, 'Nota 15, p. 79'),
     ('Receita liquida pela Nota 30', [24761056, 14437363, -966169, 1676434, -1702667, -113967], 38092050, 'Nota 30, p. 131; DRE p. 28'),
     ('Receita liquida pela nota de segmentos', [25471321, 14437363, -1816634], 38092050, 'p. 130'),
     ('Ativo total: circulante + nao circulante', [18556510, 85646257], 104202767, 'Balanco p. 24'),
     ('Passivo total: circ. + nao circ. + PL', [16552914, 45248729, 42401124], 104202767, 'Balanco pp. 26-27'),
     ('Lucro liquido: LAIR - IR e CSLL', [11701851, -3239792], 8462059, 'DRE p. 28'),
     ('Variacao de caixa: FCO + investimento + financiamento', [8361115, -15703622, 10323127], 2980620, 'DFC'),
     ('Caixa final: saldo inicial + variacao', [1682606, 2980620], 4663226, 'Balanco p. 24'),
     ('Divida liquida: emprestimos - caixa - aplicacoes', [40142347, -4663226, -7707745], 27771376, 'Nota 5, p. 55'),
     ('Provisoes judiciais liquidas por tipo', [132901, 51282, 153711, 11590, 1393402, 160526], 1903412, 'Nota 23, p. 110')]
for i, (n, comp, s, fo) in enumerate(C, 5):
    put(ws, i, 1, n)
    for j in range(6): put(ws, i, 2 + j, comp[j] if j < len(comp) else None, NUM, True)
    put(ws, i, 8, f'=SUM(B{i}:G{i})', NUM); put(ws, i, 9, s, NUM, True); put(ws, i, 10, f'=H{i}-I{i}', NUM, bold=True); put(ws, i, 11, fo)
put(ws, 19, 1, 'Total de diferencas absolutas', bold=True); put(ws, 19, 10, '=SUMPRODUCT(ABS(J5:J17))', NUM, bold=True)
widths(ws, [52] + [13] * 6 + [14, 16, 12, 30])
# ---------------- Analiticos ----------------
ws = wb.create_sheet('Analiticos'); title(ws, 'PT 520 - Procedimentos analiticos (R$ mil)', 'IPCA 2025 = 4,26% (IBGE, F6). Limite de investigacao vem da aba Materialidade.')
put(ws, 3, 1, 'IPCA 2025'); put(ws, 3, 2, 0.0426, PCT, True); put(ws, 3, 3, 'Limite de investigacao'); put(ws, 3, 4, '=Materialidade!B19', NUM)
hdr(ws, 5, ['Analitico', 'Base 1', 'Base 2', 'Registrado', 'Esperado', 'Diferenca', 'Acima do limite?', 'Metodo / premissa'])
A = [('A1. Receita bruta de saneamento', 23894855, None, 24761056, '=B6*(1+$B$3)', 'Valor de 2024 x (1 + IPCA)'),
     ('A2. Atualizacao do ativo financeiro', 17601626, 2387270, 1676434, '=B7*$B$3+C7*$B$3/2', 'Saldo inicial x IPCA + metade das transferencias x IPCA'),
     ('A3. Amortizacao do intangivel da URAE-1', 911134, None, 1765919, '=B8*2', 'Amortizacao do 2o semestre de 2024 x 2'),
     ('A4. IR e CSLL', 11701851, 48153, 3239792, '=0.34*(B9-C9)-0.34*1800000', '34% x (LAIR - equivalencia) - 34% x JCP de R$ 1,8 bi (suposicao do aluno)'),
     ('A5. PECLD', 557789, 23894855, 62067, '=B10/C10*D6', 'Percentual de 2024 sobre receita bruta x receita bruta de 2025'),
     ('A6. Despesas administrativas', 2311438, None, 1018282, '=B11*(1+$B$3)', 'Valor de 2024 x (1 + IPCA)')]
for i, (n, b1, b2, rg, ex, me) in enumerate(A, 6):
    put(ws, i, 1, n); put(ws, i, 2, b1, NUM, True); put(ws, i, 3, b2, NUM, True); put(ws, i, 4, rg, NUM, True); put(ws, i, 5, ex, NUM)
    put(ws, i, 6, f'=D{i}-E{i}', NUM); put(ws, i, 7, f'=IF(ABS(F{i})>$D$3,"Sim","Nao")'); put(ws, i, 8, me)
put(ws, 13, 1, 'Reconciliacao da PECLD (Nota 10c, p. 60): baixas - recuperacoes - reversoes (2) = despesa', bold=True)
for j, v in enumerate([501786, -401667, -29099, -8953], 2): put(ws, 14, j, v, NUM, True)
put(ws, 14, 1, 'Componentes'); put(ws, 14, 6, '=SUM(B14:E14)', NUM, bold=True); put(ws, 14, 7, '=F14-D10', NUM); put(ws, 14, 8, 'Diferenca para a PECLD registrada')
widths(ws, [44, 14, 14, 14, 14, 14, 16, 70])
# ---------------- Sensibilidade ----------------
ws = wb.create_sheet('Sensibilidade'); title(ws, 'PT 540 - Sensibilidade do ativo financeiro (R$ mil)', 'Saldo (Nota 15, F1 p. 79). Premissas variaveis em azul.')
put(ws, 4, 1, 'Saldo inicial'); put(ws, 4, 2, 17601626, NUM, True); put(ws, 5, 1, 'Transferencias do ano'); put(ws, 5, 2, 2387270, NUM, True)
put(ws, 6, 1, 'Saldo final'); put(ws, 6, 2, 21665330, NUM, True); put(ws, 7, 1, 'Base de atualizacao'); put(ws, 7, 2, '=B4+B5/2', NUM)
put(ws, 8, 1, 'Efeito divulgado da mudanca de estimativa (p. 80)'); put(ws, 8, 2, 1000000, NUM, True)
hdr(ws, 10, ['Premissa variada', 'Variacao', 'Efeito', '% da MG', '% da ME', 'Faixa'])
S = [('IPCA +/- 1 ponto percentual', 0.01, '=B7*B11'), ('Alcance da majoracao de PIS e Cofins +/- 25%', 0.25, '=B8*B12'), ('Base de ativos reversiveis +/- 5%', 0.05, '=B6*B13'), ('Base de ativos reversiveis +/- 1%', 0.01, '=B6*B14')]
for i, (n, v, f) in enumerate(S, 11):
    put(ws, i, 1, n); put(ws, i, 2, v, PCT, True); put(ws, i, 3, f, NUM); put(ws, i, 4, f'=C{i}/Materialidade!$B$12', PCT); put(ws, i, 5, f'=C{i}/Materialidade!$B$14', PCT)
    put(ws, i, 6, f'=IF(C{i}>=Materialidade!$B$14,"Acima da ME",IF(C{i}>=Materialidade!$B$12*0.5,"Entre 50% da MG e a ME","Abaixo"))')
widths(ws, [52, 12, 16, 12, 12, 26])
# ---------------- Amostragem ----------------
j = json.load(open(os.path.join(R, 'PT530_resumo.json'), encoding='utf-8'))
ws = wb.create_sheet('Amostragem'); title(ws, 'PT 530 - Amostragem por unidade monetaria (MUS), base SINTETICA', 'Parametros editaveis em azul; resultados da execucao do script 03 (valores em R$ mil).')
P_ = [('Populacao (valor liquido registrado)', j['populacao_mil'], NUM1, True), ('Distorcao toleravel (= ME)', '=Materialidade!B14', NUM1, False), ('Fator de confianca (95%, zero erros)', j['fator_confianca'], '0.00', True),
      ('Fator de expansao (95%)', j['fator_expansao'], '0.00', True), ('Distorcao esperada', j['distorcao_esperada_mil'], NUM1, True)]
for i, (n, v, f, inp) in enumerate(P_, 4): put(ws, i, 1, n); put(ws, i, 2, v, f, inp)
put(ws, 9, 1, 'Intervalo de amostragem = (TM - EM x FE) / FC', bold=True); put(ws, 9, 2, '=(B5-B8*B7)/B6', NUM1, bold=True)
put(ws, 10, 1, 'Tamanho esperado = populacao / intervalo'); put(ws, 10, 2, '=B4/B9', NUM1)
put(ws, 11, 1, 'Precisao basica = FC x intervalo'); put(ws, 11, 2, '=B6*B9', NUM1)
put(ws, 13, 1, 'Resultados da execucao (script 03; semente 241003473)', bold=True)
R_ = [('Itens-chave (valor >= intervalo)', j['itens_chave']), ('Itens por MUS', j['itens_mus']), ('Total de itens testados', j['itens_total']), ('Itens com distorcao', j['itens_com_distorcao']),
      ('Distorcao projetada (mais provavel)', j['distorcao_projetada_mil']), ('Limite superior de distorcao (Stringer)', j['limite_superior_stringer_mil']),
      ('Distorcao real inserida na base (verdade de terreno)', j['distorcao_real_inserida_mil'])]
for i, (n, v) in enumerate(R_, 14): put(ws, i, 1, n); put(ws, i, 2, v, NUM1, True)
put(ws, 22, 1, 'Limite superior < distorcao toleravel?', bold=True); put(ws, 22, 2, '=IF(B19<B5,"Sim","Nao")', bold=True)
put(ws, 23, 1, 'Distorcao projetada / distorcao real'); put(ws, 23, 2, '=B18/B20', PCT)
widths(ws, [58, 18])
# ---------------- Amostra ----------------
am = pd.read_csv(os.path.join(R, 'PT530_amostra_e_resultados.csv'))
ws = wb.create_sheet('Amostra'); title(ws, 'PT 530 - Itens selecionados e resultado do teste (SIMULADO)', 'Valor correto vem da verdade de terreno sintetica; distorcao e taint sao formulas.')
hdr(ws, 4, ['id_unidade', 'Classe', 'Tipo', 'Municipio (sintetico)', 'Selecao', 'Valor registrado (mil)', 'Valor correto (mil)', 'Distorcao (mil)', 'Taint (MUS)'])
for i, r in enumerate(am.itertuples(), 5):
    put(ws, i, 1, r.id_unidade); put(ws, i, 2, r.classe_contabil); put(ws, i, 3, r.tipo_ativo); put(ws, i, 4, r.municipio_cod); put(ws, i, 5, r.selecao)
    put(ws, i, 6, round(r.vl_registrado_mil, 3), '#,##0.000', True); put(ws, i, 7, round(r.vl_correto_mil, 3), '#,##0.000', True)
    put(ws, i, 8, f'=F{i}-G{i}', '#,##0.000'); put(ws, i, 9, f'=IF(E{i}="MUS",H{i}/F{i},"")', '0.0000')
last = 4 + len(am); put(ws, last + 2, 1, 'Total de distorcao nos itens', bold=True); put(ws, last + 2, 8, f'=SUM(H5:H{last})', '#,##0.000', bold=True)
widths(ws, [12, 38, 24, 16, 24, 18, 18, 16, 12]); ws.freeze_panes = 'A5'
# ---------------- Conc_Base ----------------
t1 = pd.read_csv(os.path.join(R, 'T01_conciliacao_nota14.csv'))
ws = wb.create_sheet('Conc_Base'); title(ws, 'T-01 - Conciliacao da base sintetica com a Nota 14 (R$ mil)', 'Base = soma da base sintetica; Nota 14 = F1, p. 72')
hdr(ws, 4, ['Classe', 'Custo - base', 'Custo - Nota 14', 'Amort. - base', 'Amort. - Nota 14', 'VL - base', 'VL - Nota 14', 'Dif. VL'])
for i, r in enumerate(t1.itertuples(), 5):
    put(ws, i, 1, r.classe)
    if r.classe == 'TOTAL':
        for c in range(2, 8): put(ws, i, c, f'=SUM({L(c)}5:{L(c)}{i-1})', NUM, bold=True)
    else:
        for c, v in zip(range(2, 8), [r.custo_base_mil, r.custo_nota14_mil, r.amort_base_mil, r.amort_nota14_mil, r.vl_base_mil, r.vl_nota14_mil]):
            put(ws, i, c, round(v, 3), '#,##0.000', True)
    put(ws, i, 8, f'=F{i}-G{i}', '#,##0.000', bold=True)
widths(ws, [44, 16, 16, 16, 16, 16, 16, 12])
# ---------------- Erros ----------------
e = pd.read_csv(os.path.join(R, 'comparacao_erros_inseridos_vs_identificados.csv'))
ws = wb.create_sheet('Erros'); title(ws, 'Erros inseridos x erros identificados pelos testes', 'Base SINTETICA; verdade de terreno em dados/registro_erros_inseridos.csv')
hdr(ws, 4, ['Codigo', 'Teste', 'Inseridas', 'Identificadas', 'Verdadeiros positivos', 'Falsos positivos', 'Falsos negativos', 'Precisao', 'Cobertura', 'Distorcao inserida (mil)', 'Distorcao identificada (mil)'])
for i, r in enumerate(e.itertuples(), 5):
    put(ws, i, 1, r.codigo); put(ws, i, 2, r.teste); put(ws, i, 3, r.inseridas, NUM, True); put(ws, i, 4, r.identificadas, NUM, True); put(ws, i, 5, r.verdadeiros_positivos, NUM, True)
    put(ws, i, 6, f'=D{i}-E{i}', NUM); put(ws, i, 7, f'=C{i}-E{i}', NUM); put(ws, i, 8, f'=E{i}/D{i}', PCT); put(ws, i, 9, f'=E{i}/C{i}', PCT)
    put(ws, i, 10, round(r.distorcao_inserida_mil, 1), NUM1, True); put(ws, i, 11, round(r.distorcao_identificada_mil, 1), NUM1, True)
put(ws, 9, 1, 'Total', bold=True)
for c in (3, 4, 5, 6, 7, 10, 11): put(ws, 9, c, f'=SUM({L(c)}5:{L(c)}8)', NUM1, bold=True)
put(ws, 11, 1, 'Nota: T-05 estagio 1 (automatico) listou 360 unidades, das quais 200 eram falsos positivos (ampliacao de capacidade apos manutencao) resolvidos pela leitura do texto.')
widths(ws, [10, 20, 11, 13, 14, 13, 13, 11, 11, 18, 20])
# ---------------- Distorcoes (PT 450) ----------------
ws = wb.create_sheet('Distorcoes'); title(ws, 'PT 450 - Mapa de distorcoes (SINTETICAS), R$ mil', 'S-01 a S-04 inseridas e confirmadas pelos testes; S-05 e julgamental, fora da base (PT 540), valor sintetico')
hdr(ws, 4, ['Cod.', 'Descricao', 'Tipo', 'Valor', 'Situacao', 'Fonte do valor'])
DD = [('S-01', 'Corte: obras de jan/2026 capitalizadas em dez/2025', 'Fato', round(e.distorcao_inserida_mil[0], 1), 'Nao corrigida', 'Base sintetica (T-04)'),
      ('S-02', 'Classificacao: manutencao capitalizada', 'Fato', round(e.distorcao_inserida_mil[1], 1), 'Nao corrigida', 'Base sintetica (T-05)'),
      ('S-03', 'Amortizacao: vida util registrada maior que a padrao', 'Fato', round(e.distorcao_inserida_mil[2], 1), 'Corrigida', 'Base sintetica (T-06)'),
      ('S-04', 'Duplicidade de unidades', 'Fato', round(e.distorcao_inserida_mil[3], 1), 'Corrigida', 'Base sintetica (T-03)'),
      ('S-05', 'Estimativa do ativo financeiro: premissa x faixa do auditor', 'Julgamento', 105000, 'Nao corrigida', 'Valor sintetico provisorio (fora da base)')]
for i, r in enumerate(DD, 5):
    for c, v in enumerate(r, 1): put(ws, i, c, v, NUM1 if c == 4 else None, c == 4)
put(ws, 11, 1, 'Corrigidas', bold=True); put(ws, 11, 4, '=SUMIF(E5:E9,"Corrigida",D5:D9)', NUM1, bold=True); put(ws, 11, 5, '=D11/Materialidade!B12', PCT)
put(ws, 12, 1, 'Nao corrigidas (agregado)', bold=True); put(ws, 12, 4, '=SUMIF(E5:E9,"Nao corrigida",D5:D9)', NUM1, bold=True); put(ws, 12, 5, '=D12/Materialidade!B12', PCT); put(ws, 12, 6, '% da MG')
put(ws, 13, 1, 'Nao corrigidas / ME'); put(ws, 13, 4, '=D12/Materialidade!B14', PCT)
put(ws, 14, 1, 'Nao corrigidas / LAIR'); put(ws, 14, 4, '=D12/Materialidade!B5', PCT)
put(ws, 15, 1, 'Itens acima do limite trivial'); put(ws, 15, 4, '=COUNTIF(D5:D9,">"&Materialidade!B16)')
widths(ws, [10, 56, 12, 14, 16, 40])
wb.save(os.path.join(P, 'Papeis_de_trabalho_calculos.xlsx')); print('ok')
